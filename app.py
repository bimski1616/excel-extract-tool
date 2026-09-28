import io
import re

import pandas as pd
import streamlit as st
import xlrd

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


TARGET_COLUMNS = [
    "Status",
    "DPP",
    "Project ID",
]


# ============================================================
# NORMALIZE TEXT
# ============================================================

def normalize_text(value):

    if value is None:
        return ""

    text = str(value)

    text = text.replace("\n", " ")

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip().casefold()


# ============================================================
# LOAD EXCEL
# ============================================================

def load_excel(file):

    file.seek(0)

    file_name = file.name.lower()

    # --------------------------------------------------------
    # XLS
    # --------------------------------------------------------

    if file_name.endswith(".xls"):

        workbook = xlrd.open_workbook(
            file_contents=file.read(),
            formatting_info=True,
        )

        return workbook, "xls"

    # --------------------------------------------------------
    # XLSX / XLSM
    # --------------------------------------------------------

    elif (
        file_name.endswith(".xlsx")
        or file_name.endswith(".xlsm")
    ):

        workbook = load_workbook(
            filename=file,
            data_only=True,
        )

        return workbook, "xlsx"

    else:

        raise ValueError(
            "Format file tidak didukung."
        )


# ============================================================
# SHEET NAMES
# ============================================================

def get_sheet_names(
    workbook,
    file_type
):

    if file_type == "xls":

        return workbook.sheet_names()

    return workbook.sheetnames


# ============================================================
# GET SHEET
# ============================================================

def get_sheet(
    workbook,
    file_type,
    sheet_name
):

    if file_type == "xls":

        return workbook.sheet_by_name(
            sheet_name
        )

    return workbook[
        sheet_name
    ]


# ============================================================
# XLS MERGED LOOKUP
# ============================================================

def build_merged_lookup_xls(sheet):

    lookup = {}

    """
    xlrd merged_cells format:

    (
        row_low,
        row_high,
        col_low,
        col_high
    )

    row_high / col_high bersifat exclusive.
    """

    for (
        row_low,
        row_high,
        col_low,
        col_high
    ) in sheet.merged_cells:

        # Value cell kiri atas
        value = sheet.cell_value(
            row_low,
            col_low
        )

        for row in range(
            row_low,
            row_high
        ):

            for col in range(
                col_low,
                col_high
            ):

                lookup[
                    (row, col)
                ] = {
                    "value": value,

                    "range": (
                        f"{get_column_letter(col_low + 1)}"
                        f"{row_low + 1}:"
                        f"{get_column_letter(col_high)}"
                        f"{row_high}"
                    ),

                    "min_row": row_low,
                    "max_row": row_high - 1,

                    "min_col": col_low,
                    "max_col": col_high - 1,
                }

    return lookup


# ============================================================
# XLSX MERGED LOOKUP
# ============================================================

def build_merged_lookup_xlsx(sheet):

    lookup = {}

    for merged_range in sheet.merged_cells.ranges:

        min_row = merged_range.min_row
        max_row = merged_range.max_row

        min_col = merged_range.min_col
        max_col = merged_range.max_col

        value = sheet.cell(
            row=min_row,
            column=min_col
        ).value

        for row in range(
            min_row,
            max_row + 1
        ):

            for col in range(
                min_col,
                max_col + 1
            ):

                lookup[
                    (row, col)
                ] = {
                    "value": value,
                    "range": str(merged_range),

                    "min_row": min_row,
                    "max_row": max_row,

                    "min_col": min_col,
                    "max_col": max_col,
                }

    return lookup


# ============================================================
# BUILD MERGED LOOKUP
# ============================================================

def build_merged_lookup(
    sheet,
    file_type
):

    if file_type == "xls":

        return build_merged_lookup_xls(
            sheet
        )

    return build_merged_lookup_xlsx(
        sheet
    )


# ============================================================
# GET VALUE
# ============================================================

def get_cell_value(
    sheet,
    file_type,
    row,
    col
):

    if file_type == "xls":

        return sheet.cell_value(
            row,
            col
        )

    return sheet.cell(
        row=row,
        column=col
    ).value


# ============================================================
# LOGICAL VALUE
# ============================================================

def get_logical_value(
    sheet,
    file_type,
    row,
    col,
    merged_lookup
):

    # Merged cell
    if (
        row,
        col
    ) in merged_lookup:

        return merged_lookup[
            (row, col)
        ]["value"]

    # Normal cell
    return get_cell_value(
        sheet,
        file_type,
        row,
        col
    )


# ============================================================
# SHEET DIMENSION
# ============================================================

def get_sheet_dimension(
    sheet,
    file_type
):

    if file_type == "xls":

        return (
            sheet.nrows,
            sheet.ncols
        )

    return (
        sheet.max_row,
        sheet.max_column
    )


# ============================================================
# FIND TARGET HEADERS
# ============================================================

def find_target_headers(
    sheet,
    file_type,
    target_columns,
    merged_lookup
):

    max_row, max_col = get_sheet_dimension(
        sheet,
        file_type
    )

    found = {}

    # XLS index mulai 0
    # XLSX index mulai 1

    if file_type == "xls":

        row_range = range(
            0,
            max_row
        )

        col_range = range(
            0,
            max_col
        )

    else:

        row_range = range(
            1,
            max_row + 1
        )

        col_range = range(
            1,
            max_col + 1
        )

    for row in row_range:

        for col in col_range:

            value = get_logical_value(
                sheet,
                file_type,
                row,
                col,
                merged_lookup
            )

            normalized_value = normalize_text(
                value
            )

            if not normalized_value:
                continue

            for target in target_columns:

                if (
                    normalize_text(target)
                    == normalized_value
                ):

                    if target not in found:

                        if file_type == "xls":

                            excel_cell = (
                                f"{get_column_letter(col + 1)}"
                                f"{row + 1}"
                            )

                        else:

                            excel_cell = (
                                f"{get_column_letter(col)}"
                                f"{row}"
                            )

                        found[target] = {
                            "row": row,
                            "col": col,
                            "cell": excel_cell,
                            "value": value,
                        }

    return found


# ============================================================
# EXTRACT DATA
# ============================================================

def extract_data(
    sheet,
    file_type,
    header_info,
    target_columns,
    merged_lookup
):

    if not header_info:

        return pd.DataFrame(
            columns=target_columns
        )

    # Header row terakhir
    data_start_row = (
        max(
            info["row"]
            for info in header_info.values()
        )
        + 1
    )

    max_row, _ = get_sheet_dimension(
        sheet,
        file_type
    )

    records = []

    for row in range(
        data_start_row,
        max_row
    ):

        record = {}

        has_value = False

        for target in target_columns:

            if target not in header_info:

                record[target] = None

                continue

            col = header_info[
                target
            ]["col"]

            value = get_logical_value(
                sheet,
                file_type,
                row,
                col,
                merged_lookup
            )

            record[target] = value

            if (
                value is not None
                and str(value).strip() != ""
            ):

                has_value = True

        if has_value:

            records.append(
                record
            )

    return pd.DataFrame(
        records,
        columns=target_columns
    )


# ============================================================
# ADD PROJECT ID 21 CHARACTERS
# ============================================================

def add_project_id_21(df):

    df = df.copy()

    if "Project ID" in df.columns:

        df["Project ID 21"] = (
            df["Project ID"]
            .astype("string")
            .str[:21]
        )

    return df


# ============================================================
# CONVERT DPP TO NUMBER
# ============================================================

def convert_dpp_to_number(value):

    if value is None:
        return None

    if pd.isna(value):
        return None

    # Kalau sudah numeric
    if isinstance(
        value,
        (int, float)
    ):

        return float(value)

    value = str(value).strip()

    if value == "":
        return None

    # --------------------------------------------------------
    # Hilangkan Rp / IDR
    # --------------------------------------------------------

    value = re.sub(
        r"(?i)rp\.?|idr",
        "",
        value
    )

    # Hilangkan spasi
    value = value.replace(
        " ",
        ""
    )

    # --------------------------------------------------------
    # Format Indonesia:
    #
    # 1.234.567,89
    #
    # menjadi:
    #
    # 1234567.89
    # --------------------------------------------------------

    if (
        "." in value
        and "," in value
        and value.rfind(",") > value.rfind(".")
    ):

        value = value.replace(
            ".",
            ""
        )

        value = value.replace(
            ",",
            "."
        )

    # --------------------------------------------------------
    # Format:
    #
    # 1,234,567
    # --------------------------------------------------------

    elif "," in value:

        value = value.replace(
            ",",
            ""
        )

    # --------------------------------------------------------
    # Format:
    #
    # 1.234.567
    # --------------------------------------------------------

    elif value.count(".") > 1:

        value = value.replace(
            ".",
            ""
        )

    # --------------------------------------------------------
    # Convert
    # --------------------------------------------------------

    try:

        return float(value)

    except (
        ValueError,
        TypeError
    ):

        return None


# ============================================================
# FORMAT DPP TO IDR DISPLAY
# ============================================================

def format_idr(value):

    if value is None or pd.isna(value):
        return "-"

    try:
        value = float(value)

        return (
            f"Rp {value:,.2f}"
            .replace(",", "X")
            .replace(".", ",")
            .replace("X", ".")
        )

    except (ValueError, TypeError):
        return "-"


# ============================================================
# CREATE SUMMARY
# ============================================================

def create_summary(df):

    # --------------------------------------------------------
    # DPP
    # --------------------------------------------------------

    summary = df.copy()

    if "DPP" in summary.columns:

        summary["DPP"] = (
            summary["DPP"]
            .apply(
                convert_dpp_to_number
            )
        )

    # --------------------------------------------------------
    # SUMMARY BERDASARKAN STATUS
    # --------------------------------------------------------

    summary = (
        summary
        .groupby(
            "Status",
            dropna=False
        )["DPP"]
        .sum()
        .reset_index()
    )

    return summary


# ============================================================
# CREATE EXCEL FILE
# ============================================================

def create_excel_file(df):

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="xlsxwriter"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="Extracted Data"
        )

    return output.getvalue()


# ============================================================
# STREAMLIT
# ============================================================

st.set_page_config(
    page_title="Excel Extract Tool",
    page_icon="📊",
    layout="wide"
)

st.title(
    "📊 Excel Extract Tool"
)

st.write(
    "Extract Status, DPP, dan Project ID "
    "dari Excel dengan merged cell."
)


# ============================================================
# UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "Upload Excel",
    type=[
        "xls",
        "xlsx",
        "xlsm"
    ]
)

if uploaded_file is None:

    st.info(
        "Upload file Excel untuk mulai."
    )

    st.stop()


# ============================================================
# LOAD
# ============================================================

try:

    workbook, file_type = load_excel(
        uploaded_file
    )

except Exception as e:

    st.error(
        f"Gagal membaca file: {e}"
    )

    st.stop()


# ============================================================
# SHEET
# ============================================================

sheet_names = get_sheet_names(
    workbook,
    file_type
)

selected_sheet = st.selectbox(
    "Pilih Sheet",
    sheet_names
)


sheet = get_sheet(
    workbook,
    file_type,
    selected_sheet
)


# ============================================================
# MERGED
# ============================================================

merged_lookup = build_merged_lookup(
    sheet,
    file_type
)


# ============================================================
# HEADER
# ============================================================

st.subheader(
    "🔎 Header Detection"
)

header_info = find_target_headers(
    sheet,
    file_type,
    TARGET_COLUMNS,
    merged_lookup
)


header_result = []

for target in TARGET_COLUMNS:

    if target in header_info:

        info = header_info[target]

        header_result.append({
            "Target": target,
            "Status": "FOUND",
            "Cell": info["cell"],
            "Value": info["value"]
        })

    else:

        header_result.append({
            "Target": target,
            "Status": "NOT FOUND",
            "Cell": "-",
            "Value": "-"
        })


st.dataframe(
    pd.DataFrame(header_result),
    use_container_width=True,
    hide_index=True
)


# ============================================================
# EXTRACT
# ============================================================

missing = [
    target
    for target in TARGET_COLUMNS
    if target not in header_info
]


if missing:

    st.warning(
        "Header tidak ditemukan: "
        + ", ".join(missing)
    )


if st.button(
    "🚀 Extract Data",
    type="primary",
    use_container_width=True
):

    if missing:

        st.error(
            "Masih ada header yang belum ditemukan."
        )

    else:

        result_df = extract_data(
            sheet,
            file_type,
            header_info,
            TARGET_COLUMNS,
            merged_lookup
        )

        result_df = result_df.dropna(
            how="all"
        ).reset_index(
            drop=True
        )

        # ====================================================
        # AMBIL 21 KARAKTER PERTAMA PROJECT ID
        # ====================================================

        result_df = add_project_id_21(
            result_df
        )

        # ====================================================
        # FORMAT DPP MENJADI NUMBER
        # ====================================================

        if "DPP" in result_df.columns:

            result_df["DPP"] = (
                result_df["DPP"]
                .apply(
                    convert_dpp_to_number
                )
            )

        # ====================================================
        # SAVE RESULT
        # ====================================================

        st.session_state[
            "result_df"
        ] = result_df


# ============================================================
# RESULT
# ============================================================

if "result_df" in st.session_state:

    result_df = st.session_state[
        "result_df"
    ]

    st.subheader(
        "📊 Result"
    )

    st.dataframe(
        result_df.style.format(
            {
                "DPP": format_idr
            }
        ),
        use_container_width=True
    )


    # ========================================================
    # SUMMARY PREVIEW
    # ========================================================

    st.subheader(
        "📋 Summary Preview"
    )

    summary_df = create_summary(
        result_df
    )

    st.dataframe(
        summary_df.style.format(
            {
                "DPP": format_idr
            }
        ),
        use_container_width=True,
        hide_index=True
    )


    # ========================================================
    # DOWNLOAD
    # ========================================================

    st.subheader(
        "💾 Download"
    )

    download_format = st.radio(
        "Pilih format file:",
        [
            "CSV",
            "Excel"
        ],
        horizontal=True
    )


    # ========================================================
    # DOWNLOAD CSV
    # ========================================================

    if download_format == "CSV":

        csv_data = result_df.to_csv(
            index=False
        ).encode(
            "utf-8-sig"
        )

        st.download_button(
            label="⬇️ Download CSV",
            data=csv_data,
            file_name="extracted_data.csv",
            mime="text/csv",
            use_container_width=True
        )


    # ========================================================
    # DOWNLOAD EXCEL
    # ========================================================

    elif download_format == "Excel":

        excel_data = create_excel_file(
            result_df
        )

        st.download_button(
            label="⬇️ Download Excel",
            data=excel_data,
            file_name="extracted_data.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True
        )